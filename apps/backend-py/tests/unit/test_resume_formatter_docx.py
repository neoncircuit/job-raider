"""
Tests for DOCX resume output.

Before 2026-10-18, ``ResumeFormatter._format_docx`` passed ``bold=`` and
``italic=`` to python-docx ``Paragraph.add_run``, which accepts only text
and style. Every resume with experience, projects, or education raised
``TypeError``. ``format_resume`` caught it, so PDF output still worked and
the DOCX file was silently missing.
"""

from pathlib import Path

import pytest

from src.generation.formatter import DOCX_AVAILABLE, ResumeFormatter
from src.generation.resume_writer import GeneratedResume

pytestmark = pytest.mark.skipif(not DOCX_AVAILABLE, reason="python-docx missing")


def make_resume() -> GeneratedResume:
    """Build a resume with every section that uses styled runs.

    Returns:
        A generated resume with experience, projects, and education.
    """
    return GeneratedResume(
        summary="Backend engineer focused on Python services.",
        skills=["Python", "FastAPI", "PostgreSQL"],
        experience=[
            {
                "title": "Software Engineer",
                "company": "Example Pte Ltd",
                "dates": "2024 - 2026",
                "highlights": ["Built a job ingestion service."],
            }
        ],
        projects=[
            {
                "name": "Job Raider",
                "description": "Local-first job search pipeline.",
                "technologies": ["Python", "Next.js"],
                "highlights": ["Deterministic cover-letter grounding."],
            }
        ],
        education=[{"degree": "BSc Computer Science", "school": "NUS", "year": "2024"}],
        raw_response="{}",
        model_used="test-model",
    )


def test_docx_export_succeeds_with_styled_sections(tmp_path: Path) -> None:
    """DOCX export writes the file and styles titles, dates, and labels.

    Args:
        tmp_path: Temporary output directory.
    """
    from docx import Document

    result = ResumeFormatter(output_dir=str(tmp_path)).format_resume(
        make_resume(), filename="resume_test", formats=["docx"]
    )

    assert result.errors == []
    assert result.docx_path is not None
    assert Path(result.docx_path).exists()

    runs = [run for p in Document(result.docx_path).paragraphs for run in p.runs]
    bold = {run.text for run in runs if run.bold}
    italic = {run.text.strip() for run in runs if run.italic}

    assert {"Software Engineer", "Technologies: ", "BSc Computer Science"} <= bold
    assert "2024 - 2026" in italic
